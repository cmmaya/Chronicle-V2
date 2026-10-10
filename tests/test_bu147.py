"""BU147: text extraction for session documents (.txt, .md, .pdf, .docx)."""
import pytest

from src.assistant.session_documents import (
    DocumentError, extract_document, is_allowed_document,
)
from src.config import SESSION_DOCUMENTS

SYLLABUS = 'Week 1 covers sampling methods. Week 2 covers the survey critique.'


def _text_pdf(path, lines):
    """A one-page PDF with a text layer (hand-built; no PDF writer needed)."""
    content = 'BT /F1 12 Tf 72 720 Td ' + ' T* '.join(
        f'({line}) Tj' for line in lines) + ' ET'
    objects = [
        '<< /Type /Catalog /Pages 2 0 R >>',
        '<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        '<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] '
        '/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>',
        f'<< /Length {len(content)} >>\nstream\n{content}\nendstream',
        '<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
    ]
    out, offsets = b'%PDF-1.4\n', []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f'{number} 0 obj\n{body}\nendobj\n'.encode('latin-1')
    xref = len(out)
    out += f'xref\n0 {len(objects) + 1}\n0000000000 65535 f \n'.encode()
    out += ''.join(f'{o:010d} 00000 n \n' for o in offsets).encode()
    out += (f'trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n'
            f'startxref\n{xref}\n%%EOF\n').encode()
    path.write_bytes(out)
    return path


def _blank_pdf(path):
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(612, 792)
    with open(path, 'wb') as handle:
        writer.write(handle)
    return path


def _docx(path, paragraphs=(), table=None, header=None):
    from docx import Document
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    if header:
        document.sections[0].header.paragraphs[0].text = header
    document.save(str(path))
    return path


def test_txt_and_md_are_read(tmp_path):
    for suffix in ('txt', 'md'):
        path = tmp_path / f'notes.{suffix}'
        path.write_text(f'# Notes\r\n\r\n\r\n\r\n{SYLLABUS}   \r\n', encoding='utf-8')
        doc = extract_document(str(path))
        assert (doc.name, doc.file_type) == (f'notes.{suffix}', suffix)
        assert doc.text == f'# Notes\n\n{SYLLABUS}'
        assert doc.char_count == len(doc.text)
        assert doc.file_bytes == path.stat().st_size


def test_txt_in_cp1252_decodes(tmp_path):
    path = tmp_path / 'cafe.txt'
    path.write_bytes(('Café opens at nine. ' * 4).encode('cp1252'))
    assert 'Café opens' in extract_document(str(path)).text


def test_pdf_pages_are_joined(tmp_path):
    doc = extract_document(str(_text_pdf(tmp_path / 's.pdf', [SYLLABUS])))
    assert doc.file_type == 'pdf'
    assert 'sampling methods' in doc.text and 'survey critique' in doc.text


def test_docx_reads_paragraphs_then_table_rows_and_skips_headers(tmp_path):
    path = _docx(tmp_path / 'h.docx', [SYLLABUS], table=[['Week', 'Topic'], ['1', 'Sampling']],
                 header='CONFIDENTIAL HEADER')
    doc = extract_document(str(path))
    assert doc.text == f'{SYLLABUS}\n\nWeek | Topic\n1 | Sampling'.replace('\n\nWeek', '\nWeek')
    assert 'CONFIDENTIAL' not in doc.text


def test_refused_inputs(tmp_path, monkeypatch):
    with pytest.raises(DocumentError, match='no longer exists'):
        extract_document(str(tmp_path / 'missing.txt'))
    with pytest.raises(DocumentError, match='Folders'):
        extract_document(str(tmp_path))
    odd = tmp_path / 'slides.pptx'
    odd.write_bytes(b'x' * 100)
    with pytest.raises(DocumentError, match='not supported'):
        extract_document(str(odd))
    big = tmp_path / 'big.txt'
    big.write_text('word ' * 100, encoding='utf-8')
    monkeypatch.setitem(SESSION_DOCUMENTS, 'max_file_bytes', 100)
    with pytest.raises(DocumentError, match='the limit is'):
        extract_document(str(big))


def test_binary_behind_a_txt_name_is_refused(tmp_path):
    path = tmp_path / 'blob.txt'
    path.write_bytes(b'abc\x00' * 50)
    with pytest.raises(DocumentError, match='not readable text'):
        extract_document(str(path))


def test_scanned_pdf_and_empty_docx_have_no_readable_text(tmp_path):
    with pytest.raises(DocumentError, match='scanned PDFs are not supported'):
        extract_document(str(_blank_pdf(tmp_path / 'scan.pdf')))
    with pytest.raises(DocumentError, match='No readable text found in empty.docx$'):
        extract_document(str(_docx(tmp_path / 'empty.docx')))


def test_corrupt_files_raise_document_error(tmp_path):
    for name in ('bad.pdf', 'bad.docx'):
        path = tmp_path / name
        path.write_bytes(b'this is not really a document' * 5)
        with pytest.raises(DocumentError, match=f'Could not read {name}'):
            extract_document(str(path))


def test_encrypted_pdf_is_refused(tmp_path):
    from pypdf import PdfReader, PdfWriter
    writer = PdfWriter(clone_from=PdfReader(str(_text_pdf(tmp_path / 'p.pdf', [SYLLABUS]))))
    writer.encrypt('secret')
    locked = tmp_path / 'locked.pdf'
    with open(locked, 'wb') as handle:
        writer.write(handle)
    with pytest.raises(DocumentError, match='password-protected'):
        extract_document(str(locked))


def test_over_max_text_chars_is_refused_not_truncated(tmp_path, monkeypatch):
    path = tmp_path / 'long.txt'
    path.write_text('word ' * 100, encoding='utf-8')
    monkeypatch.setitem(SESSION_DOCUMENTS, 'max_text_chars', 100)
    with pytest.raises(DocumentError, match='has 499 characters'):
        extract_document(str(path))


def test_is_allowed_document(tmp_path):
    (tmp_path / 'a.PDF').write_bytes(b'x')
    assert is_allowed_document(str(tmp_path / 'a.PDF'))
    assert not is_allowed_document(str(tmp_path / 'a.pptx'))
    assert not is_allowed_document(str(tmp_path))
    assert not is_allowed_document('')
