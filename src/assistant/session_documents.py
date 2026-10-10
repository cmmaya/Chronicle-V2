"""Turn a document into the plain text a session stores (BU147).

``extract_document`` is the one place a ``.txt``, ``.md``, ``.pdf`` or
``.docx`` file becomes text. Qt-free and network-free; synchronous, so callers
run it on a background job.

A file that yields no usable text is refused here rather than stored as an
empty document that would silently answer nothing - notably a scanned PDF,
which has no text layer (there is no OCR).
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass
from types import SimpleNamespace
from typing import List, Tuple

from ..config import SESSION_DOCUMENTS
from .live_qa import strip_transcript_evidence
from .reference_doc import ReferenceDocError, decode_text, select_relevant_excerpt


class DocumentError(Exception):
    """A document could not be read. Carries a message fit to show a user."""


@dataclass
class ExtractedDocument:
    name: str
    file_type: str  # txt | md | pdf | docx
    file_bytes: int
    text: str
    char_count: int


def is_allowed_document(path: str) -> bool:
    """Whether ``path`` is a file with an allowed extension (for drag-enter)."""
    if not path or os.path.isdir(path):
        return False
    return os.path.splitext(path)[1].lower() in SESSION_DOCUMENTS["allowed_extensions"]


def extract_document(path: str) -> ExtractedDocument:
    """Read ``path`` into an ``ExtractedDocument``.

    Raises ``DocumentError`` for anything that does not yield usable text -
    never a partial result and never a library exception.
    """
    name = os.path.basename(path or "")
    if not path or not os.path.exists(path):
        raise DocumentError(f"{name or 'That file'} no longer exists")
    if os.path.isdir(path):
        raise DocumentError("Folders cannot be added - drop a file")

    extension = os.path.splitext(path)[1].lower()
    allowed = SESSION_DOCUMENTS["allowed_extensions"]
    if extension not in allowed:
        raise DocumentError(
            f"{name} is not supported. Use {', '.join(allowed)} files")

    try:
        size = os.path.getsize(path)
    except OSError as e:
        raise DocumentError(f"Could not read {name}: {e}") from e
    cap = SESSION_DOCUMENTS["max_file_bytes"]
    if size > cap:
        raise DocumentError(
            f"{name} is {_megabytes(size)}; the limit is {_megabytes(cap)}")

    file_type = extension[1:]
    try:
        if file_type in ("txt", "md"):
            raw_text = _read_plain(path, name)
        elif file_type == "pdf":
            raw_text = _read_pdf(path, name)
        else:
            raw_text = _read_docx(path)
    except DocumentError:
        raise
    except Exception as e:  # a library's own error types are not the UI's business
        raise DocumentError(f"Could not read {name}: {e}") from e

    text = _normalise(raw_text)
    if len(text) < SESSION_DOCUMENTS["min_text_chars"]:
        hint = " - scanned PDFs are not supported" if file_type == "pdf" else ""
        raise DocumentError(f"No readable text found in {name}{hint}")
    limit = SESSION_DOCUMENTS["max_text_chars"]
    if len(text) > limit:
        raise DocumentError(
            f"{name} has {len(text):,} characters of text; the limit is {limit:,}")
    return ExtractedDocument(
        name=name, file_type=file_type, file_bytes=size, text=text,
        char_count=len(text))


def is_allowed_image(path: str) -> bool:
    """Whether ``path`` is an image a session can take as context."""
    if not path or os.path.isdir(path):
        return False
    return os.path.splitext(path)[1].lower() in SESSION_DOCUMENTS["image_extensions"]


def image_document_text(name: str, summary: str, visible_text: List[str]) -> str:
    """The text an image is stored as: its description, then its words."""
    parts = [f"Image: {name}", "", "Description:", summary.strip() or "(none)"]
    lines = [line.strip() for line in visible_text if line and line.strip()]
    if lines:
        parts.extend(["", "Visible text:", *lines])
    return "\n".join(parts)


def extract_image(path: str, describe=None) -> ExtractedDocument:
    """Turn an image into an ``ExtractedDocument`` through a vision model.

    ``describe(path) -> {"summary", "visible_text"}`` defaults to
    ``ScreenshotContextGenerator.describe_image``; it makes a network call, so
    run this on a background job. Raises ``DocumentError`` on any failure.
    """
    name = os.path.basename(path or "")
    if not path or not os.path.isfile(path):
        raise DocumentError(f"{name or 'That file'} no longer exists")
    extension = os.path.splitext(path)[1].lower()
    if extension not in SESSION_DOCUMENTS["image_extensions"]:
        raise DocumentError(f"{name} is not a supported image")
    try:
        size = os.path.getsize(path)
    except OSError as e:
        raise DocumentError(f"Could not read {name}: {e}") from e
    cap = SESSION_DOCUMENTS["max_image_bytes"]
    if size > cap:
        raise DocumentError(
            f"{name} is {_megabytes(size)}; the limit is {_megabytes(cap)}")

    if describe is None:
        from ..screenshots.context_generator import ScreenshotContextGenerator
        describe = ScreenshotContextGenerator().describe_image
    try:
        context = describe(path)
    except Exception as e:  # network, key or model errors
        raise DocumentError(f"Could not describe {name}: {e}") from e

    summary = str(context.get("summary") or "")
    visible = list(context.get("visible_text") or [])
    if not summary.strip() and not visible:
        raise DocumentError(f"Nothing could be read from {name}")
    text = image_document_text(name, summary, visible)
    return ExtractedDocument(
        name=name, file_type=extension[1:], file_bytes=size, text=text,
        char_count=len(text))


def is_allowed_attachment(path: str) -> bool:
    """Whether ``path`` can be added to a session, as a document or an image."""
    return is_allowed_document(path) or is_allowed_image(path)


def is_image_type(file_type: str) -> bool:
    """Whether a stored document's ``file_type`` came from an image."""
    return f".{(file_type or '').lower()}" in SESSION_DOCUMENTS["image_extensions"]


def attachment_extensions() -> List[str]:
    return SESSION_DOCUMENTS["allowed_extensions"] + SESSION_DOCUMENTS["image_extensions"]


def extract_attachment(path: str) -> ExtractedDocument:
    """``extract_image`` for images, ``extract_document`` for everything else."""
    return extract_image(path) if is_allowed_image(path) else extract_document(path)


def _megabytes(size: int) -> str:
    return f"{size / (1024 * 1024):.1f} MB"


def _read_plain(path: str, name: str) -> str:
    with open(path, "rb") as handle:
        raw = handle.read()
    # The extension proves nothing about the contents.
    if b"\x00" in raw:
        raise DocumentError(f"{name} is not readable text")
    try:
        return decode_text(raw)
    except ReferenceDocError as e:
        raise DocumentError(f"{name}: {e}") from e


def _read_pdf(path: str, name: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    if reader.is_encrypted and not reader.decrypt(""):
        raise DocumentError("This PDF is password-protected")
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    return "\n\n".join(page for page in pages if page)


def _read_docx(path: str) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(path)
    blocks: List[str] = []
    # Body order, so a table sits where the author put it. Headers and footers
    # live outside the body and are ignored.
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            blocks.append(Paragraph(child, document).text)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, document).rows:
                cells, seen = [], set()
                for cell in row.cells:
                    if id(cell._tc) in seen:  # a merged cell repeats
                        continue
                    seen.add(id(cell._tc))
                    if cell.text.strip():
                        cells.append(cell.text.strip())
                if cells:
                    blocks.append(" | ".join(cells))
    return "\n".join(blocks)


def _normalise(text: str) -> str:
    """One paragraph convention for every type: ``\\n`` line ends, no trailing
    spaces, at most one blank line in a row."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# --- Rendering for the Specific Session prompt (BU149) -----------------------

WHOLE = "whole"
EXCERPT = "excerpt"


def document_header(document: dict) -> str:
    return f"[D{document['id']}] {document['name']} ({document['file_type']})"


def render_documents(documents: List[dict], question: str) -> Tuple[str, str]:
    """``(mode, block)`` for a session's documents, in the order added.

    ``whole``: when the documents together fit ``whole_total_max_chars``, each
    is sent in full - the same for every question, so it can sit in the cached
    part of the prompt. Otherwise ``excerpt``: each document contributes the
    paragraphs that best match ``question`` (and nothing else of the
    transcript evidence a live-answer question carries); a document with no
    excerpt is listed by header only so the model knows it exists.
    """
    total = sum(len(d["text"]) for d in documents)
    if total <= SESSION_DOCUMENTS["whole_total_max_chars"]:
        parts = ["## Session documents"]
        for document in documents:
            parts.extend(["", document_header(document), document["text"]])
        return WHOLE, "\n".join(parts)

    bare = strip_transcript_evidence(question)
    cap = SESSION_DOCUMENTS["excerpt_chars_per_document"]
    parts = ["## Session document excerpts"]
    for document in documents:
        excerpt = select_relevant_excerpt(
            SimpleNamespace(text=document["text"]), bare, max_chars=cap)
        parts.extend(["", document_header(document)])
        if excerpt:
            parts.append(excerpt)
    return EXCERPT, "\n".join(parts)
